import streamlit as st
import fitz  
from docx import Document
from transformers import T5Tokenizer, T5ForConditionalGeneration
from fpdf import FPDF
import tempfile

MODEL_NAME = "Timosh-nlp/nlp_10"  
tokenizer = T5Tokenizer.from_pretrained(MODEL_NAME)
model = T5ForConditionalGeneration.from_pretrained(MODEL_NAME)

def extract_text_from_pdf(pdf_file):
    doc = fitz.open(stream=pdf_file.read(), filetype="pdf")
    text = ""
    for page in doc:
        text += page.get_text("text") + "\n"
    return text

def extract_text_from_docx(docx_file):
    doc = Document(docx_file)
    text = "\n".join([para.text for para in doc.paragraphs])
    return text

def generate_questions(text, num_questions):
    input_text = "generate question: " + text
    input_ids = tokenizer.encode(input_text, return_tensors="pt", max_length=512, truncation=True)
    
    outputs = model.generate(input_ids, max_length=128, num_return_sequences=num_questions, do_sample=True)
    
    questions = [tokenizer.decode(output, skip_special_tokens=True) for output in outputs]
    return questions

def generate_text_file(questions):
    text_content = "\n".join([f"{i+1}. {q}" for i, q in enumerate(questions)])
    return text_content

def generate_pdf_file(questions):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, "Generated Questions", ln=True, align="C")
    pdf.ln(10)
    
    for i, q in enumerate(questions, 1):
        pdf.multi_cell(0, 10, f"{i}. {q}")
        pdf.ln(5)
    
    temp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    pdf.output(temp_pdf.name)
    return temp_pdf.name

def main():
    st.title("📄 Document Question Generator")
    st.write("Upload a **PDF** or **Word Document**, and generate questions.")

    uploaded_file = st.file_uploader("Upload your file", type=["pdf", "docx"])
    num_questions = st.number_input("Number of Questions to Generate", min_value=1, max_value=20, value=5)

    if uploaded_file is not None:
        file_extension = uploaded_file.name.split(".")[-1].lower()

        if file_extension == "pdf":
            text = extract_text_from_pdf(uploaded_file)
        elif file_extension == "docx":
            text = extract_text_from_docx(uploaded_file)
        else:
            st.error("Unsupported file format!")
            return

        if st.button("Generate Questions"):
            with st.spinner("Generating questions... ⏳"):
                questions = generate_questions(text, num_questions)
            
            # Display questions in a table
            st.write("### Generated Questions")
            st.table({"Question Number": list(range(1, len(questions) + 1)), "Question": questions})
            
            # Provide download options
            text_content = generate_text_file(questions)
            text_file = tempfile.NamedTemporaryFile(delete=False, suffix=".txt")
            with open(text_file.name, "w") as f:
                f.write(text_content)
            
            pdf_file_path = generate_pdf_file(questions)
            
            st.download_button("📥 Download as TXT", data=text_content, file_name="questions.txt", mime="text/plain")
            with open(pdf_file_path, "rb") as pdf_file:
                st.download_button("📥 Download as PDF", data=pdf_file, file_name="questions.pdf", mime="application/pdf")

if __name__ == "__main__":
    main()
